from PyQt6.QtCore import QThread, pyqtSignal
import pandas as pd
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo

class SaveFileThread(QThread):
    finished = pyqtSignal(str)
    error = pyqtSignal(str)

    def __init__(self, df: pd.DataFrame, output_file: str, operation_name: str, dual_sheet: bool = False,
                 df2: pd.DataFrame = None, sheet1_name: str = None, sheet2_name: str = None):
        super().__init__()
        self.output_file = output_file
        self.operation_name = operation_name
        self.df = df
        self.df2 = df2
        self.dual_sheet = dual_sheet
        self.sheet1_name = sheet1_name
        self.sheet2_name = sheet2_name

    def run(self):
        try:
            if not self.dual_sheet:
                self.df.to_excel(self.output_file, index=False, sheet_name=self.operation_name)
            else:
                with pd.ExcelWriter(self.output_file, engine='openpyxl') as writer:
                    self.df.to_excel(writer, index=False, sheet_name=self.sheet1_name)
                    self.df2.to_excel(writer, index=False, sheet_name=self.sheet2_name)

            # Open and adjust the workbook
            wb = openpyxl.load_workbook(self.output_file)
            if not self.dual_sheet:
                ws = wb[self.operation_name]
                ws.add_table(self.create_table(self.df, self.operation_name))
                self.adjust_columns(ws)
            else:
                ws = wb[self.sheet1_name]
                ws.add_table(self.create_table(self.df, self.sheet1_name))
                self.adjust_columns(ws)
                ws = wb[self.sheet2_name]
                ws.add_table(self.create_table(self.df2, self.sheet2_name))
                self.adjust_columns(ws)

            wb.save(self.output_file)
            wb.close()
            self.finished.emit(self.output_file)
        except Exception as e:
            self.error.emit(str(e))

    @staticmethod
    def create_table(df: pd.DataFrame, display_name: str):
        table_ref = f'A1:{chr(64 + len(df.columns))}{len(df) + 1}'
        table = Table(displayName=display_name, ref=table_ref)
        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )
        table.tableStyleInfo = style
        return table

    @staticmethod
    def adjust_columns(ws: openpyxl.worksheet):
        min_width = 8
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            header_width = len(str(column[0].value)) + 8

            # Find the maximum length of the values in the column
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except (AttributeError, TypeError):
                    pass

            # Set the column width to the maximum length + a buffer (e.g. 2), but ensure it's at least min_width
            ws.column_dimensions[column_letter].width = max(header_width, max_length + 2, min_width)
        return ws
